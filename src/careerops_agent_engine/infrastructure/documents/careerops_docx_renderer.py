"""Deterministic DOCX rendering for the standard CareerOps CV template."""

from io import BytesIO
from typing import cast
from zipfile import ZipFile, ZipInfo

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from docx.styles.style import ParagraphStyle
from docx.text.paragraph import Paragraph

from careerops_agent_engine.application.exceptions import (
    CVRenderingError,
)
from careerops_agent_engine.application.ports.cv_renderer import (
    CVTemplateRenderer,
    RenderedCVDocument,
)
from careerops_agent_engine.domain.enums import (
    CVArtifactFormat,
)
from careerops_agent_engine.domain.models.cv_version import (
    CVVersion,
)
from careerops_agent_engine.domain.models.structured_cv import (
    CVEntry,
    StructuredCVSection,
)

DOCX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)

CANONICAL_ZIP_TIMESTAMP = (
    1980,
    1,
    1,
    0,
    0,
    0,
)


class CareerOpsStandardDocxRenderer(CVTemplateRenderer):
    """Render an ATS-friendly one-column CareerOps DOCX document."""

    TEMPLATE_ID = "careerops-standard"
    TEMPLATE_VERSION = "1.0.0"

    @property
    def template_id(self) -> str:
        """Return the standard template identifier."""

        return self.TEMPLATE_ID

    @property
    def template_version(self) -> str:
        """Return the standard template version."""

        return self.TEMPLATE_VERSION

    def render(
        self,
        *,
        version: CVVersion,
    ) -> RenderedCVDocument:
        """Render one immutable structured CV to deterministic DOCX bytes."""

        self._validate_template_provenance(version)

        document = Document()

        self._configure_document(document)

        self._render_preamble(
            document,
            version,
        )

        for section in version.structured_cv.sections:
            self._render_section(
                document,
                section,
            )

        stream = BytesIO()

        try:
            document.save(stream)
        except Exception as exc:
            raise CVRenderingError("DOCX rendering failed.") from exc

        canonical_bytes = canonicalize_docx_package(stream.getvalue())

        return RenderedCVDocument(
            artifact_format=(CVArtifactFormat.DOCX),
            media_type=DOCX_MEDIA_TYPE,
            filename=(f"{version.cv_version_id}.docx"),
            data=canonical_bytes,
        )

    def _validate_template_provenance(
        self,
        version: CVVersion,
    ) -> None:
        """Require version provenance to select this exact renderer."""

        if version.provenance.template_id != self.template_id:
            raise CVRenderingError(
                "CV version template identifier does not match the renderer."
            )

        if version.provenance.template_version != self.template_version:
            raise CVRenderingError(
                "CV version template version does not match the renderer."
            )

    @staticmethod
    def _configure_document(
        document: DocumentObject,
    ) -> None:
        """Apply deterministic A4 and typography defaults."""

        section = document.sections[0]

        section.page_width = Mm(210)
        section.page_height = Mm(297)

        section.top_margin = Mm(14)
        section.bottom_margin = Mm(14)
        section.left_margin = Mm(16)
        section.right_margin = Mm(16)

        normal = cast(
            ParagraphStyle,
            document.styles["Normal"],
        )

        normal.font.name = "Arial"
        normal.font.size = Pt(9.5)

        normal.paragraph_format.space_after = Pt(2)

        normal.paragraph_format.line_spacing = 1.0

        core = document.core_properties

        core.title = "CareerOps CV"
        core.author = "CareerOps"
        core.subject = "Evidence-grounded tailored CV"
        core.comments = "Generated deterministically from an approved CV version."

    @staticmethod
    def _render_preamble(
        document: DocumentObject,
        version: CVVersion,
    ) -> None:
        """Render source preamble without changing its text."""

        preamble = version.structured_cv.preamble_text

        if preamble is None:
            return

        lines = preamble.split("\n")

        for index, line in enumerate(lines):
            paragraph = document.add_paragraph()

            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

            paragraph.paragraph_format.space_after = Pt(1)

            run = paragraph.add_run(line)

            if index == 0:
                run.bold = True
                run.font.size = Pt(17)
            else:
                run.font.size = Pt(9)

    def _render_section(
        self,
        document: DocumentObject,
        section: StructuredCVSection,
    ) -> None:
        """Render one structured section in source order."""

        heading = document.add_paragraph()

        heading.paragraph_format.space_before = Pt(7)

        heading.paragraph_format.space_after = Pt(3)

        run = heading.add_run(section.heading.upper())

        run.bold = True
        run.font.size = Pt(10.5)

        self._add_bottom_border(heading)

        if section.free_text is not None:
            self._render_free_text(
                document,
                section.free_text,
            )

        for entry in section.entries:
            self._render_entry(
                document,
                entry,
            )

    @staticmethod
    def _render_free_text(
        document: DocumentObject,
        text: str,
    ) -> None:
        """Render preserved source text without paraphrasing it."""

        for line in text.split("\n"):
            paragraph = document.add_paragraph()

            paragraph.paragraph_format.space_after = Pt(1.5)

            paragraph.add_run(line)

    @staticmethod
    def _render_entry(
        document: DocumentObject,
        entry: CVEntry,
    ) -> None:
        """Render one explicitly structured CV entry."""

        title = document.add_paragraph()

        title.paragraph_format.space_before = Pt(3)

        title.paragraph_format.space_after = Pt(1)

        title_run = title.add_run(entry.title)

        title_run.bold = True
        title_run.font.size = Pt(10)

        metadata = [
            value
            for value in (
                entry.subtitle,
                entry.location,
                entry.date_text,
            )
            if value is not None
        ]

        if metadata:
            metadata_paragraph = document.add_paragraph(" | ".join(metadata))

            metadata_paragraph.paragraph_format.space_after = Pt(1.5)

            metadata_run = metadata_paragraph.runs[0]

            metadata_run.italic = True
            metadata_run.font.size = Pt(9)

        for bullet in entry.bullets:
            paragraph = document.add_paragraph(style="List Bullet")

            paragraph.paragraph_format.space_after = Pt(1)

            paragraph.add_run(bullet.text)

    @staticmethod
    def _add_bottom_border(
        paragraph: Paragraph,
    ) -> None:
        """Add a simple section divider without layout tables."""

        properties = paragraph._p.get_or_add_pPr()

        borders = properties.find(qn("w:pBdr"))

        if borders is None:
            borders = OxmlElement("w:pBdr")

            properties.append(borders)

        bottom = OxmlElement("w:bottom")

        bottom.set(
            qn("w:val"),
            "single",
        )

        bottom.set(
            qn("w:sz"),
            "6",
        )

        bottom.set(
            qn("w:space"),
            "1",
        )

        bottom.set(
            qn("w:color"),
            "808080",
        )

        borders.append(bottom)


def canonicalize_docx_package(
    data: bytes,
) -> bytes:
    """Normalize ZIP metadata so identical content yields identical bytes."""

    source_stream = BytesIO(data)

    target_stream = BytesIO()

    try:
        with (
            ZipFile(
                source_stream,
                "r",
            ) as source,
            ZipFile(
                target_stream,
                "w",
            ) as target,
        ):
            for source_info in sorted(
                source.infolist(),
                key=lambda item: item.filename,
            ):
                target_info = ZipInfo(
                    source_info.filename,
                    date_time=(CANONICAL_ZIP_TIMESTAMP),
                )

                target_info.compress_type = source_info.compress_type

                target_info.external_attr = source_info.external_attr

                target_info.create_system = source_info.create_system

                target.writestr(
                    target_info,
                    source.read(source_info.filename),
                )

    except Exception as exc:
        raise CVRenderingError("Rendered DOCX package is invalid.") from exc

    return target_stream.getvalue()
