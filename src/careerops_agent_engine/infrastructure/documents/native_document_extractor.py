"""Native PDF, DOCX and UTF-8 text extraction."""

from io import BytesIO
from zipfile import BadZipFile

from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from careerops_agent_engine.application.exceptions import (
    DocumentExtractionError,
    DocumentTextUnavailableError,
)
from careerops_agent_engine.domain.enums import (
    CareerDocumentFormat,
)
from careerops_agent_engine.domain.models.document import (
    ExtractedDocumentText,
)


class NativeDocumentExtractor:
    """Extract text from validated PDF, DOCX and UTF-8 bytes."""

    def extract(
        self,
        *,
        document_id: str,
        document_format: CareerDocumentFormat,
        data: bytes,
    ) -> ExtractedDocumentText:
        """Extract text using the format-specific native parser."""

        if document_format is CareerDocumentFormat.PDF:
            return extract_pdf(
                document_id=document_id,
                data=data,
            )

        if document_format is CareerDocumentFormat.DOCX:
            return extract_docx(
                document_id=document_id,
                data=data,
            )

        if document_format is CareerDocumentFormat.TEXT:
            return extract_text(
                document_id=document_id,
                data=data,
            )

        raise DocumentExtractionError("Unsupported career-document format.")


def extract_pdf(
    *,
    document_id: str,
    data: bytes,
) -> ExtractedDocumentText:
    """Extract native text from a validated PDF."""

    try:
        reader = PdfReader(
            BytesIO(data),
            strict=False,
        )
    except (
        PdfReadError,
        OSError,
        ValueError,
    ) as exc:
        raise DocumentExtractionError("The PDF could not be parsed.") from exc

    page_texts: list[str] = []
    warnings: list[str] = []

    for page_number, page in enumerate(
        reader.pages,
        start=1,
    ):
        try:
            raw_text = page.extract_text()
        except Exception as exc:
            raise DocumentExtractionError(
                f"PDF text extraction failed on page {page_number}."
            ) from exc

        cleaned = normalise_text(raw_text or "")

        if not cleaned:
            warnings.append(f"Page {page_number} contained no extractable native text.")
            continue

        page_texts.append(cleaned)

    text = "\n\n".join(page_texts).strip()

    if not text:
        raise DocumentTextUnavailableError(
            "The PDF contains no extractable native text."
        )

    return ExtractedDocumentText(
        document_id=document_id,
        text=text,
        page_count=len(reader.pages),
        paragraph_count=None,
        warnings=warnings,
    )


def extract_docx(
    *,
    document_id: str,
    data: bytes,
) -> ExtractedDocumentText:
    """Extract paragraph and table text from a DOCX package."""

    try:
        document = Document(BytesIO(data))
    except (
        BadZipFile,
        PackageNotFoundError,
        OSError,
        ValueError,
        KeyError,
    ) as exc:
        raise DocumentExtractionError("The DOCX document could not be parsed.") from exc

    text_blocks: list[str] = []

    paragraph_count = 0

    for paragraph in document.paragraphs:
        cleaned = normalise_text(paragraph.text)

        if not cleaned:
            continue

        text_blocks.append(cleaned)
        paragraph_count += 1

    for table in document.tables:
        for row in table.rows:
            cells = [normalise_text(cell.text) for cell in row.cells]

            populated_cells = [cell for cell in cells if cell]

            if populated_cells:
                text_blocks.append(" | ".join(populated_cells))

    text = "\n".join(text_blocks).strip()

    if not text:
        raise DocumentTextUnavailableError(
            "The DOCX document contains no extractable native text."
        )

    return ExtractedDocumentText(
        document_id=document_id,
        text=text,
        page_count=None,
        paragraph_count=paragraph_count,
        warnings=[],
    )


def extract_text(
    *,
    document_id: str,
    data: bytes,
) -> ExtractedDocumentText:
    """Extract and normalise one trusted UTF-8 text source."""

    try:
        raw_text = data.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise DocumentExtractionError("The text source is not valid UTF-8.") from exc

    text = normalise_text(raw_text)

    if not text:
        raise DocumentTextUnavailableError("The text source contains no usable text.")

    return ExtractedDocumentText(
        document_id=document_id,
        text=text,
        page_count=None,
        paragraph_count=len(text.splitlines()),
        warnings=[],
    )


def normalise_text(
    text: str,
) -> str:
    """Normalise extracted whitespace without rewriting content."""

    lines = [" ".join(line.split()) for line in text.splitlines()]

    populated_lines = [line for line in lines if line]

    return "\n".join(populated_lines)
