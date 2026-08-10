"""Deterministic verification of CareerOps standard DOCX output."""

from io import BytesIO
from zipfile import BadZipFile, ZipFile

from docx import Document

from careerops_agent_engine.application.ports.cv_artifact_verifier import (
    CVArtifactVerificationResult,
    CVArtifactVerifier,
)
from careerops_agent_engine.domain.enums import (
    CVArtifactFormat,
)
from careerops_agent_engine.domain.models.cv_version import (
    CVVersion,
)
from careerops_agent_engine.domain.models.structured_cv import (
    CVEntry,
    StructuredCVSection,
)


class CareerOpsStandardDocxVerifier(CVArtifactVerifier):
    """Verify stored standard-template DOCX content."""

    @property
    def artifact_format(self) -> CVArtifactFormat:
        """Return the supported generated-artifact format."""

        return CVArtifactFormat.DOCX

    def verify(
        self,
        *,
        version: CVVersion,
        data: bytes,
    ) -> CVArtifactVerificationResult:
        """Verify package validity and visible CV content."""

        if not data:
            return failed("Stored DOCX bytes are empty.")

        package_result = self._verify_package(data)

        if package_result is not None:
            return package_result

        try:
            document = Document(BytesIO(data))

        except Exception:
            return failed("Stored DOCX could not be opened.")

        actual_paragraphs = [paragraph.text for paragraph in document.paragraphs]

        expected_paragraphs = build_expected_visible_paragraphs(version)

        if actual_paragraphs != expected_paragraphs:
            return failed(
                "DOCX visible content does not match the structured CV version."
            )

        visible_text = "\n".join(actual_paragraphs)

        internal_ids = collect_internal_provenance_ids(version)

        if any(identifier in visible_text for identifier in internal_ids):
            return failed("DOCX exposes internal CareerOps provenance identifiers.")

        if document.tables:
            return failed(
                "The standard CareerOps DOCX must remain "
                "a one-column document without layout tables."
            )

        return CVArtifactVerificationResult(passed=True)

    @staticmethod
    def _verify_package(
        data: bytes,
    ) -> CVArtifactVerificationResult | None:
        """Validate the underlying Open XML ZIP package."""

        try:
            with ZipFile(
                BytesIO(data),
                "r",
            ) as package:
                names = set(package.namelist())

                if (
                    "[Content_Types].xml" not in names
                    or "word/document.xml" not in names
                ):
                    return failed(
                        "Stored DOCX package is missing required Open XML files."
                    )

                corrupted_member = package.testzip()

                if corrupted_member is not None:
                    return failed("Stored DOCX ZIP package is corrupted.")

        except BadZipFile:
            return failed("Stored DOCX package is invalid.")

        return None


def build_expected_visible_paragraphs(
    version: CVVersion,
) -> list[str]:
    """Build the exact paragraph contract of the standard renderer."""

    expected: list[str] = []

    preamble = version.structured_cv.preamble_text

    if preamble is not None:
        expected.extend(preamble.split("\n"))

    for section in version.structured_cv.sections:
        expected.extend(expected_section_paragraphs(section))

    return expected


def expected_section_paragraphs(
    section: StructuredCVSection,
) -> list[str]:
    """Build visible paragraphs for one rendered section."""

    paragraphs = [section.heading.upper()]

    if section.free_text is not None:
        paragraphs.extend(section.free_text.split("\n"))

    for entry in section.entries:
        paragraphs.extend(expected_entry_paragraphs(entry))

    return paragraphs


def expected_entry_paragraphs(
    entry: CVEntry,
) -> list[str]:
    """Build visible paragraphs for one structured CV entry."""

    paragraphs = [entry.title]

    metadata = [
        value
        for value in (
            entry.subtitle,
            entry.location,
            entry.date_text,
        )
        if value is not None
    ]

    if metadata:
        paragraphs.append(" | ".join(metadata))

    paragraphs.extend(bullet.text for bullet in entry.bullets)

    return paragraphs


def collect_internal_provenance_ids(
    version: CVVersion,
) -> set[str]:
    """Collect internal identifiers that must never appear visibly."""

    identifiers = {
        version.cv_version_id,
        version.provenance.job_id,
        version.provenance.thread_id,
        *version.provenance.requirement_ids,
        *version.provenance.supporting_evidence_ids,
        *version.provenance.final_proposal_ids,
    }

    for change in version.applied_changes:
        identifiers.add(change.change_id)

    return {identifier for identifier in identifiers if identifier}


def failed(
    note: str,
) -> CVArtifactVerificationResult:
    """Create one failed deterministic verification result."""

    return CVArtifactVerificationResult(
        passed=False,
        notes=(note,),
    )
