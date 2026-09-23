"""Run a real PostgreSQL and configured-LLM evidence persistence smoke test."""

from io import BytesIO
from uuid import uuid4

from docx import Document
from langchain_core.tracers.langchain import wait_for_all_tracers
from sqlalchemy import text

from careerops_agent_engine.api.dependencies import (
    get_career_document_repository,
    get_cv_document_ingestion_service,
    get_cv_document_preparation_service,
    get_cv_document_upload_service,
    get_cv_evidence_audit_repository,
    get_cv_evidence_workflow_service,
    get_database_engine,
    get_database_session_factory,
    get_document_storage,
    get_evidence_repository,
)
from careerops_agent_engine.domain.enums import (
    CVEvidenceReviewRunStatus,
)
from careerops_agent_engine.domain.models.evidence_review import (
    EvidenceReviewDecision,
)

DOCX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)


def build_synthetic_cv_docx() -> bytes:
    """Create a small native-text DOCX with deterministic CV headings."""

    document = Document()

    document.add_paragraph("CareerOps 7G4C Smoke Candidate")

    document.add_heading("Profile", level=1)
    document.add_paragraph(
        "AI engineering graduate building production-style Python applications."
    )

    document.add_heading("Skills", level=1)
    document.add_paragraph(
        "Python, FastAPI, Docker, PostgreSQL, LangGraph, and LangSmith."
    )

    document.add_heading("Projects", level=1)
    document.add_paragraph("CareerOps Agent Engine")
    document.add_paragraph(
        "Built a stateful career workflow using Python, FastAPI, LangGraph, "
        "PostgreSQL, Docker, and LangSmith."
    )
    document.add_paragraph(
        "Added human approval before career evidence is promoted to the trusted "
        "evidence registry."
    )

    document.add_heading("Education", level=1)
    document.add_paragraph("BSc Computer Science, University of Southampton.")

    stream = BytesIO()
    document.save(stream)

    return stream.getvalue()


def reset_runtime_dependencies() -> None:
    """Simulate a fresh application process with new DB-backed adapters."""

    get_database_engine().dispose()

    for provider in (
        get_cv_evidence_workflow_service,
        get_cv_document_preparation_service,
        get_cv_document_ingestion_service,
        get_cv_document_upload_service,
        get_cv_evidence_audit_repository,
        get_career_document_repository,
        get_evidence_repository,
        get_database_session_factory,
        get_database_engine,
        get_document_storage,
    ):
        provider.cache_clear()


def assert_awaiting_row(
    *,
    review_run_id: str,
) -> None:
    """Verify the PostgreSQL awaiting row uses real SQL NULL."""

    statement = text(
        """
        SELECT
            status,
            review_result IS NULL AS review_result_is_sql_null
        FROM cv_evidence_review_runs
        WHERE review_run_id = :review_run_id
        """
    )

    with get_database_engine().connect() as connection:
        row = (
            connection.execute(
                statement,
                {"review_run_id": review_run_id},
            )
            .mappings()
            .one()
        )

    assert row["status"] == "awaiting_review"
    assert bool(row["review_result_is_sql_null"]) is True

    print("PostgreSQL awaiting row verified: review_result is real SQL NULL.")


def assert_completed_database_state(
    *,
    user_id: str,
    review_run_id: str,
    expected_evidence_count: int,
) -> None:
    """Verify completed run, history and approved evidence rows."""

    with get_database_engine().connect() as connection:
        run = (
            connection.execute(
                text(
                    """
                SELECT
                    status,
                    review_result IS NOT NULL AS has_review_result
                FROM cv_evidence_review_runs
                WHERE review_run_id = :review_run_id
                """
                ),
                {"review_run_id": review_run_id},
            )
            .mappings()
            .one()
        )

        history_count = connection.execute(
            text(
                """
                SELECT COUNT(*)
                FROM cv_evidence_review_history
                WHERE review_run_id = :review_run_id
                """
            ),
            {"review_run_id": review_run_id},
        ).scalar_one()

        evidence_count = connection.execute(
            text(
                """
                SELECT COUNT(*)
                FROM career_evidence
                WHERE user_id = :user_id
                  AND verification_status = 'approved'
                """
            ),
            {"user_id": user_id},
        ).scalar_one()

    assert run["status"] == "completed"
    assert bool(run["has_review_result"]) is True
    assert history_count == 1
    assert evidence_count == expected_evidence_count

    print("PostgreSQL completed state verified.")
    print(f"  review history rows: {history_count}")
    print(f"  approved evidence rows: {evidence_count}")


def main() -> None:
    """Exercise the complete persisted CV evidence workflow."""

    user_id = f"USER-LIVE-CV-7G4C-{uuid4().hex[:8].upper()}"

    print("=== CareerOps 7G4C live smoke test ===")
    print(f"user_id: {user_id}")

    cv_bytes = build_synthetic_cv_docx()

    ingestion_service = get_cv_document_ingestion_service()

    document = ingestion_service.ingest(
        user_id=user_id,
        original_filename="careerops-live-smoke.docx",
        declared_media_type=DOCX_MEDIA_TYPE,
        data=cv_bytes,
    )

    print(f"document_id: {document.document_id}")
    print("Secure upload + PostgreSQL document metadata: PASSED")

    workflow = get_cv_evidence_workflow_service()

    awaiting = workflow.start_review(
        user_id=user_id,
        document_id=document.document_id,
    )

    assert awaiting.status is CVEvidenceReviewRunStatus.AWAITING_REVIEW, (
        "Expected the configured model to produce at least one grounded "
        "evidence proposal, "
        f"but workflow status was {awaiting.status.value!r}."
    )

    assert awaiting.proposals

    print(f"review_run_id: {awaiting.review_run_id}")
    print(f"Grounded proposals: {len(awaiting.proposals)}")
    print(f"overlap findings: {len(awaiting.overlap_findings)}")

    for proposal in awaiting.proposals:
        print(
            "  proposal:",
            proposal.proposal_id,
            "|",
            proposal.category.value,
            "|",
            proposal.title,
        )

    assert_awaiting_row(
        review_run_id=awaiting.review_run_id,
    )

    original_snapshot_json = awaiting.model_dump_json()

    print("\n--- Simulating application restart ---")
    reset_runtime_dependencies()

    restarted_workflow = get_cv_evidence_workflow_service()

    recovered = restarted_workflow.start_review(
        user_id=user_id,
        document_id=document.document_id,
    )

    assert recovered.review_run_id == awaiting.review_run_id
    assert recovered.model_dump_json() == original_snapshot_json

    print("Restart recovery: PASSED")
    print("Exact persisted review snapshot recovered.")

    proposal_ids = [proposal.proposal_id for proposal in recovered.proposals]

    acknowledged_overlap_ids = sorted(
        {finding.proposal_id for finding in recovered.overlap_findings}
    )

    decision = EvidenceReviewDecision(
        approved_proposal_ids=proposal_ids,
        acknowledged_overlap_proposal_ids=acknowledged_overlap_ids,
        reviewer_comment=(
            "7G4C live smoke test: reviewed grounded proposals "
            "and explicitly approved them."
        ),
    )

    completed = restarted_workflow.submit_review(
        user_id=user_id,
        review_run_id=recovered.review_run_id,
        decision=decision,
    )

    assert completed.status is CVEvidenceReviewRunStatus.COMPLETED
    assert completed.review_result is not None

    approved_evidence = completed.review_result.approved_evidence

    assert len(approved_evidence) == len(proposal_ids)

    print("\nHuman approval transaction: PASSED")
    print(f"approved evidence items: {len(approved_evidence)}")

    evidence_repository = get_evidence_repository()

    for evidence in approved_evidence:
        persisted = evidence_repository.get_approved(
            user_id=user_id,
            evidence_id=evidence.evidence_id,
        )

        assert persisted == evidence

        print(
            "  trusted evidence:",
            evidence.evidence_id,
            "|",
            evidence.title,
        )

    print("Trusted EvidenceRepository retrieval: PASSED")

    assert_completed_database_state(
        user_id=user_id,
        review_run_id=completed.review_run_id,
        expected_evidence_count=len(approved_evidence),
    )

    print("\n--- Simulating second restart + exact retry ---")
    reset_runtime_dependencies()

    retry_workflow = get_cv_evidence_workflow_service()

    retry_result = retry_workflow.submit_review(
        user_id=user_id,
        review_run_id=completed.review_run_id,
        decision=decision,
    )

    assert retry_result == completed

    reviews = get_cv_evidence_audit_repository().list_reviews(
        user_id=user_id,
        review_run_id=completed.review_run_id,
    )

    assert len(reviews) == 1

    print("Exact completed retry is idempotent: PASSED")
    print("Audit history remains one row.")

    print("\n=== 7G4C LIVE SMOKE TEST PASSED ===")
    print("The rows and stored DOCX are intentionally left in the")
    print("development environment so they can be inspected before cleanup.")
    print(f"user_id: {user_id}")
    print(f"document_id: {document.document_id}")
    print(f"review_run_id: {completed.review_run_id}")


if __name__ == "__main__":
    try:
        main()
    finally:
        wait_for_all_tracers()
