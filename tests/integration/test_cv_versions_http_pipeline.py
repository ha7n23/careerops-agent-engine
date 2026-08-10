"""Live HTTP proof of final CV generation, retrieval, and downloads."""

import os
from hashlib import sha256
from io import BytesIO
from pathlib import Path
from uuid import uuid4

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfReader
from sqlalchemy import delete, select

from careerops_agent_engine.api.dependencies import (
    get_artifact_storage,
    get_career_document_repository,
    get_cv_artifact_retrieval_service,
    get_cv_document_preparation_service,
    get_cv_version_repository,
    get_database_engine,
    get_database_session_factory,
    get_document_storage,
    get_docx_rendering_service,
    get_docx_verification_service,
    get_evidence_repository,
    get_final_cv_assembly_service,
    get_final_cv_generation_service,
    get_job_analysis_audit_repository,
    get_pdf_conversion_service,
    get_pdf_verification_service,
)
from careerops_agent_engine.core.config import get_settings
from careerops_agent_engine.domain.enums import (
    ApprovalStatus,
    CareerDocumentFormat,
    CareerDocumentStatus,
    CVArtifactVerificationStatus,
    CVSection,
    CVVersionStatus,
    EvidenceCategory,
    EvidenceSourceType,
    JobAnalysisRunStatus,
    VerificationStatus,
)
from careerops_agent_engine.domain.models.cv import (
    CVChangeProposal,
)
from careerops_agent_engine.domain.models.evidence import (
    SourceReference,
)
from careerops_agent_engine.infrastructure.database.models.cv_evidence import (
    CareerDocumentRecord,
)
from careerops_agent_engine.infrastructure.database.models.cv_version import (
    CVArtifactRecord,
    CVVersionRecord,
)
from careerops_agent_engine.infrastructure.database.models.evidence import (
    CareerEvidenceRecord,
)
from careerops_agent_engine.infrastructure.database.models.job_analysis import (
    JobAnalysisRunRecord,
)
from careerops_agent_engine.main import app

pytestmark = pytest.mark.integration


@pytest.mark.skipif(
    os.getenv("RUN_LIVE_DOCUMENT_TESTS") != "true",
    reason="Live document integration tests are disabled.",
)
def test_real_http_final_cv_generation_and_download_pipeline(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Prove the final CV workflow through the real FastAPI boundary."""

    suffix = uuid4().hex[:8].upper()

    user_id = f"USER-LIVE-HTTP-{suffix}"
    other_user_id = f"USER-LIVE-OTHER-{suffix}"

    document_id = f"DOC-LIVE-HTTP-{suffix}"
    thread_id = f"THR-LIVE-HTTP-{suffix}"
    job_id = f"JOB-LIVE-HTTP-{suffix}"

    proposal_id = f"CVP-LIVE-HTTP-{suffix}"
    requirement_id = f"REQ-LIVE-HTTP-{suffix}"
    evidence_id = f"EVD-LIVE-HTTP-{suffix}"

    original_text = "Built CareerOps using Python."

    final_text = "Built CareerOps using Python and FastAPI."

    # Keep this live proof isolated from normal development
    # document/artifact storage.
    monkeypatch.setenv(
        "CAREEROPS_DOCUMENT_STORAGE_ROOT",
        str(tmp_path / "documents"),
    )

    monkeypatch.setenv(
        "CAREEROPS_ARTIFACT_STORAGE_ROOT",
        str(tmp_path / "artifacts"),
    )

    # main.py loads Settings while importing the FastAPI app.
    # Clear cached runtime dependencies so the temporary roots
    # above are used by the real production providers.
    clear_final_cv_runtime_caches()

    settings = get_settings()

    engine = get_database_engine()

    session_factory = get_database_session_factory()

    source_bytes = build_source_docx(
        original_text=original_text,
    )

    storage_key = get_document_storage().save(
        user_id=user_id,
        document_id=document_id,
        document_format=(CareerDocumentFormat.DOCX),
        data=source_bytes,
    )

    proposal = CVChangeProposal(
        proposal_id=proposal_id,
        section=CVSection.PROJECTS,
        current_text=original_text,
        proposed_text=final_text,
        requirement_ids=[
            requirement_id,
        ],
        supporting_evidence_ids=[
            evidence_id,
        ],
        confidence_score=1.0,
    )

    source_reference = SourceReference(
        source_type=(EvidenceSourceType.UPLOADED_CV),
        source_id=document_id,
        source_excerpt=original_text,
    )

    try:
        # ---------------------------------------------------------
        # 1. Seed the already-approved prerequisite business state.
        #
        # We deliberately do not call an LLM in this test.
        # The HTTP pipeline begins from a completed, human-approved
        # job-analysis result and approved evidence.
        # ---------------------------------------------------------
        with session_factory.begin() as session:
            session.add(
                CareerDocumentRecord(
                    user_id=user_id,
                    document_id=document_id,
                    original_filename=("careerops-live-http-source.docx"),
                    document_format=(CareerDocumentFormat.DOCX.value),
                    media_type=(
                        "application/vnd.openxmlformats-officedocument."
                        "wordprocessingml.document"
                    ),
                    size_bytes=len(source_bytes),
                    sha256_hex=sha256(source_bytes).hexdigest(),
                    storage_key=storage_key,
                    status=(CareerDocumentStatus.EXTRACTED.value),
                )
            )

            session.add(
                CareerEvidenceRecord(
                    user_id=user_id,
                    evidence_id=evidence_id,
                    category=(EvidenceCategory.PROJECT.value),
                    title=("CareerOps Agent Engine"),
                    verification_status=(VerificationStatus.APPROVED.value),
                    technologies=[
                        "Python",
                        "FastAPI",
                    ],
                    capabilities=[
                        "AI application engineering",
                    ],
                    approved_claims=[
                        final_text,
                    ],
                    source_references=[source_reference.model_dump(mode="json")],
                )
            )

            session.add(
                JobAnalysisRunRecord(
                    thread_id=thread_id,
                    user_id=user_id,
                    job_id=job_id,
                    status=(JobAnalysisRunStatus.COMPLETED.value),
                    role_title=("Junior AI Engineer"),
                    fit_score=100.0,
                    review_status=(ApprovalStatus.APPROVED.value),
                    cv_proposals=[proposal.model_dump(mode="json")],
                    claim_verification_reports=[],
                    reviewable_proposal_ids=[proposal_id],
                    blocked_proposal_ids=[],
                    final_cv_proposals=[proposal.model_dump(mode="json")],
                )
            )

        # ---------------------------------------------------------
        # 2. Enter exclusively through the real FastAPI boundary.
        # ---------------------------------------------------------
        with TestClient(app) as client:
            request_body = {
                "thread_id": thread_id,
                "source_document_id": (document_id),
            }

            headers = {"X-User-ID": user_id}

            generated = client.post(
                "/api/v1/cv-versions",
                headers=headers,
                json=request_body,
            )

            assert generated.status_code == 200, generated.text

            generated_body = generated.json()

            cv_version_id = str(generated_body["cv_version_id"])

            assert generated_body["status"] == CVVersionStatus.VERIFIED.value

            assert generated_body["version_number"] == 1

            assert generated_body["source_document_id"] == document_id

            assert generated_body["thread_id"] == thread_id

            assert generated_body["job_id"] == job_id

            assert generated_body["reused_existing_version"] is False

            assert len(generated_body["artifacts"]) == 2

            assert all(
                artifact["verification_status"]
                == (CVArtifactVerificationStatus.VERIFIED.value)
                for artifact in generated_body["artifacts"]
            )

            # Private/backend-only provenance must not leak.
            assert "storage_key" not in generated.text

            assert evidence_id not in generated.text

            assert requirement_id not in generated.text

            assert proposal_id not in generated.text

            artifacts_by_format = {
                artifact["artifact_format"]: artifact
                for artifact in generated_body["artifacts"]
            }

            # -----------------------------------------------------
            # 3. Retrieve authoritative version metadata.
            # -----------------------------------------------------
            retrieved = client.get(
                (f"/api/v1/cv-versions/{cv_version_id}"),
                headers=headers,
            )

            assert retrieved.status_code == 200, retrieved.text

            assert retrieved.json()["cv_version_id"] == cv_version_id

            assert retrieved.json()["status"] == CVVersionStatus.VERIFIED.value

            # -----------------------------------------------------
            # 4. Download and reopen the actual generated DOCX.
            # -----------------------------------------------------
            docx_response = client.get(
                (f"/api/v1/cv-versions/{cv_version_id}/artifacts/docx"),
                headers=headers,
            )

            assert docx_response.status_code == 200, docx_response.text

            assert docx_response.content.startswith(b"PK")

            assert docx_response.headers["content-disposition"] == (
                f'attachment; filename="{cv_version_id}.docx"'
            )

            assert (
                sha256(docx_response.content).hexdigest()
                == (artifacts_by_format["docx"]["sha256_hex"])
            )

            assert (
                len(docx_response.content)
                == (artifacts_by_format["docx"]["size_bytes"])
            )

            reopened_docx = Document(BytesIO(docx_response.content))

            docx_text = "\n".join(
                paragraph.text for paragraph in reopened_docx.paragraphs
            )

            assert final_text in docx_text

            assert requirement_id not in docx_text

            assert evidence_id not in docx_text

            assert proposal_id not in docx_text

            # -----------------------------------------------------
            # 5. Download and reopen the actual LibreOffice PDF.
            # -----------------------------------------------------
            pdf_response = client.get(
                (f"/api/v1/cv-versions/{cv_version_id}/artifacts/pdf"),
                headers=headers,
            )

            assert pdf_response.status_code == 200, pdf_response.text

            assert pdf_response.content.startswith(b"%PDF-")

            assert pdf_response.headers["content-type"].startswith("application/pdf")

            assert (
                sha256(pdf_response.content).hexdigest()
                == (artifacts_by_format["pdf"]["sha256_hex"])
            )

            assert (
                len(pdf_response.content) == (artifacts_by_format["pdf"]["size_bytes"])
            )

            reader = PdfReader(BytesIO(pdf_response.content))

            assert len(reader.pages) >= 1

            pdf_text = "\n".join(page.extract_text() or "" for page in reader.pages)

            # LibreOffice/PDF extraction may introduce line
            # wrapping. Compare the approved sentence after
            # whitespace normalisation.
            normalized_pdf_text = " ".join(pdf_text.split())

            assert final_text in normalized_pdf_text

            assert requirement_id not in pdf_text

            assert evidence_id not in pdf_text

            assert proposal_id not in pdf_text

            # -----------------------------------------------------
            # 6. Exact retry must reuse the existing version.
            #    LibreOffice must not create another PDF.
            # -----------------------------------------------------
            retried = client.post(
                "/api/v1/cv-versions",
                headers=headers,
                json=request_body,
            )

            assert retried.status_code == 200, retried.text

            assert retried.json()["cv_version_id"] == cv_version_id

            assert retried.json()["reused_existing_version"] is True

            assert retried.json()["artifacts"] == generated_body["artifacts"]

            # -----------------------------------------------------
            # 7. Prove authenticated user isolation through HTTP.
            # -----------------------------------------------------
            cross_user_version = client.get(
                (f"/api/v1/cv-versions/{cv_version_id}"),
                headers={"X-User-ID": (other_user_id)},
            )

            assert cross_user_version.status_code == 404

            cross_user_docx = client.get(
                (f"/api/v1/cv-versions/{cv_version_id}/artifacts/docx"),
                headers={"X-User-ID": (other_user_id)},
            )

            assert cross_user_docx.status_code == 404

            cross_user_pdf = client.get(
                (f"/api/v1/cv-versions/{cv_version_id}/artifacts/pdf"),
                headers={"X-User-ID": (other_user_id)},
            )

            assert cross_user_pdf.status_code == 404

        # ---------------------------------------------------------
        # 8. Confirm retry created no duplicate persistent state.
        # ---------------------------------------------------------
        with session_factory() as session:
            versions = session.scalars(
                select(CVVersionRecord).where(
                    CVVersionRecord.user_id == user_id,
                    CVVersionRecord.thread_id == thread_id,
                )
            ).all()

            assert len(versions) == 1

            assert versions[0].status == (CVVersionStatus.VERIFIED.value)

            artifacts = session.scalars(
                select(CVArtifactRecord).where(
                    CVArtifactRecord.cv_version_id == (versions[0].cv_version_id)
                )
            ).all()

            assert len(artifacts) == 2

            assert all(
                artifact.verification_status
                == (CVArtifactVerificationStatus.VERIFIED.value)
                for artifact in artifacts
            )

        print()

        print("=== CareerOps live HTTP final-CV proof ===")

        print(f"user_id: {user_id}")

        print(f"cv_version_id: {cv_version_id}")

        print("POST: generated verified DOCX + PDF")

        print("GET: version metadata retrieved")

        print("DOWNLOAD: real DOCX + PDF reopened successfully")

        print("RETRY: existing CV version reused")

        print("ISOLATION: cross-user retrieval blocked")

        print("REAL FINAL-CV HTTP PIPELINE PASSED")

    finally:
        # ---------------------------------------------------------
        # Live database cleanup. Generated files live beneath
        # tmp_path and are removed automatically by pytest.
        # ---------------------------------------------------------
        with session_factory.begin() as session:
            version_ids = session.scalars(
                select(CVVersionRecord.cv_version_id).where(
                    CVVersionRecord.user_id == user_id,
                    CVVersionRecord.thread_id == thread_id,
                )
            ).all()

            if version_ids:
                session.execute(
                    delete(CVArtifactRecord).where(
                        CVArtifactRecord.cv_version_id.in_(version_ids)
                    )
                )

                session.execute(
                    delete(CVVersionRecord).where(
                        CVVersionRecord.cv_version_id.in_(version_ids)
                    )
                )

            session.execute(
                delete(JobAnalysisRunRecord).where(
                    JobAnalysisRunRecord.thread_id == thread_id
                )
            )

            session.execute(
                delete(CareerEvidenceRecord).where(
                    CareerEvidenceRecord.user_id == user_id,
                    CareerEvidenceRecord.evidence_id == evidence_id,
                )
            )

            session.execute(
                delete(CareerDocumentRecord).where(
                    CareerDocumentRecord.user_id == user_id,
                    CareerDocumentRecord.document_id == document_id,
                )
            )

        engine.dispose()

        clear_final_cv_runtime_caches()

    # Prove the HTTP request used the isolated test roots,
    # rather than the normal .careerops_data paths.
    assert settings.document_storage_root == (tmp_path / "documents")

    assert settings.artifact_storage_root == (tmp_path / "artifacts")


def build_source_docx(
    *,
    original_text: str,
) -> bytes:
    """Create a native source CV containing one unique project anchor."""

    document = Document()

    document.add_paragraph("CareerOps Test Candidate")

    document.add_paragraph("candidate@example.com")

    document.add_paragraph("Profile")

    document.add_paragraph("AI engineer building evidence-grounded LLM applications.")

    document.add_paragraph("Projects")

    document.add_paragraph("CareerOps Agent Engine")

    document.add_paragraph(original_text)

    stream = BytesIO()

    document.save(stream)

    return stream.getvalue()


def clear_final_cv_runtime_caches() -> None:
    """Clear cached dependencies affected by final-CV storage settings."""

    for provider in (
        get_final_cv_generation_service,
        get_cv_artifact_retrieval_service,
        get_final_cv_assembly_service,
        get_docx_rendering_service,
        get_docx_verification_service,
        get_pdf_conversion_service,
        get_pdf_verification_service,
        get_cv_document_preparation_service,
        get_cv_version_repository,
        get_career_document_repository,
        get_evidence_repository,
        get_job_analysis_audit_repository,
        get_document_storage,
        get_artifact_storage,
        get_database_session_factory,
        get_database_engine,
    ):
        provider.cache_clear()

    get_settings.cache_clear()
