"""Tests for native PDF and DOCX text extraction."""

from io import BytesIO

import pytest
from docx import Document
from pypdf import PdfWriter

from careerops_agent_engine.application.exceptions import (
    DocumentExtractionError,
    DocumentTextUnavailableError,
)
from careerops_agent_engine.domain.enums import (
    CareerDocumentFormat,
)
from careerops_agent_engine.infrastructure.documents.native_document_extractor import (
    NativeDocumentExtractor,
)


def build_docx_bytes() -> bytes:
    """Create a real DOCX with paragraphs and a table."""

    document = Document()

    document.add_heading(
        "Candidate Name",
        level=1,
    )

    document.add_paragraph("Python Engineer")

    table = document.add_table(
        rows=1,
        cols=2,
    )

    table.rows[0].cells[0].text = "FastAPI"

    table.rows[0].cells[1].text = "Docker"

    output = BytesIO()

    document.save(output)

    return output.getvalue()


def build_blank_pdf_bytes() -> bytes:
    """Create a structurally valid PDF with no native text."""

    writer = PdfWriter()

    writer.add_blank_page(
        width=612,
        height=792,
    )

    output = BytesIO()

    writer.write(output)

    return output.getvalue()


def build_text_pdf_bytes() -> bytes:
    """Create a minimal valid PDF containing native text."""

    content = b"BT /F1 12 Tf 72 720 Td (Python Engineer) Tj ET"

    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        (b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>"),
        (
            b"<< /Type /Page "
            b"/Parent 2 0 R "
            b"/MediaBox [0 0 612 792] "
            b"/Resources << "
            b"/Font << /F1 4 0 R >> "
            b">> "
            b"/Contents 5 0 R >>"
        ),
        (b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"),
        (
            b"<< /Length "
            + str(len(content)).encode()
            + b" >>\nstream\n"
            + content
            + b"\nendstream"
        ),
    ]

    output = bytearray(b"%PDF-1.4\n")

    offsets = [0]

    for object_number, pdf_object in enumerate(
        objects,
        start=1,
    ):
        offsets.append(len(output))

        output.extend(f"{object_number} 0 obj\n".encode())
        output.extend(pdf_object)
        output.extend(b"\nendobj\n")

    xref_offset = len(output)

    output.extend(f"xref\n0 {len(objects) + 1}\n".encode())
    output.extend(b"0000000000 65535 f \n")

    for offset in offsets[1:]:
        output.extend(f"{offset:010d} 00000 n \n".encode())

    output.extend(
        (
            "trailer\n"
            f"<< /Size {len(objects) + 1} "
            "/Root 1 0 R >>\n"
            "startxref\n"
            f"{xref_offset}\n"
            "%%EOF\n"
        ).encode()
    )

    return bytes(output)


def test_docx_extracts_paragraph_and_table_text() -> None:
    """Native DOCX text should preserve useful document content."""

    extractor = NativeDocumentExtractor()

    result = extractor.extract(
        document_id="DOC-001",
        document_format=CareerDocumentFormat.DOCX,
        data=build_docx_bytes(),
    )

    assert result.document_id == "DOC-001"
    assert "Candidate Name" in result.text
    assert "Python Engineer" in result.text
    assert "FastAPI | Docker" in result.text
    assert result.page_count is None
    assert result.paragraph_count == 2


def test_blank_pdf_requires_future_ocr_path() -> None:
    """A PDF with no native text should not silently succeed."""

    extractor = NativeDocumentExtractor()

    with pytest.raises(
        DocumentTextUnavailableError,
        match="no extractable native text",
    ):
        extractor.extract(
            document_id="DOC-001",
            document_format=CareerDocumentFormat.PDF,
            data=build_blank_pdf_bytes(),
        )


def test_invalid_pdf_bytes_raise_stable_error() -> None:
    """Parser failures should not leak pypdf exceptions."""

    extractor = NativeDocumentExtractor()

    with pytest.raises(
        DocumentExtractionError,
        match="PDF could not be parsed",
    ):
        extractor.extract(
            document_id="DOC-001",
            document_format=CareerDocumentFormat.PDF,
            data=b"%PDF-1.7\nnot actually a valid PDF",
        )


def test_invalid_docx_bytes_raise_stable_error() -> None:
    """Parser failures should not leak python-docx exceptions."""

    extractor = NativeDocumentExtractor()

    with pytest.raises(
        DocumentExtractionError,
        match="DOCX document could not be parsed",
    ):
        extractor.extract(
            document_id="DOC-001",
            document_format=CareerDocumentFormat.DOCX,
            data=b"not a docx package",
        )


def test_pdf_extracts_native_text() -> None:
    """A text-bearing PDF should produce native extracted text."""

    extractor = NativeDocumentExtractor()

    result = extractor.extract(
        document_id="DOC-001",
        document_format=CareerDocumentFormat.PDF,
        data=build_text_pdf_bytes(),
    )

    assert result.document_id == "DOC-001"
    assert "Python Engineer" in result.text
    assert result.page_count == 1
    assert result.paragraph_count is None
    assert result.warnings == []
